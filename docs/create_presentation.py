import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_documind_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    PRIMARY_NAVY = RGBColor(15, 23, 42)      # #0F172A
    SECONDARY_BLUE = RGBColor(30, 58, 138)   # #1E3A8A
    ACCENT_INDIGO = RGBColor(79, 70, 229)    # #4F46E5
    TEXT_DARK = RGBColor(30, 41, 59)         # #1E293B
    TEXT_MUTED = RGBColor(100, 116, 139)     # #64748B
    BG_LIGHT = RGBColor(248, 250, 252)       # #F8FAFC
    CARD_BG = RGBColor(255, 255, 255)        # #FFFFFF
    BORDER_COLOR = RGBColor(226, 232, 240)   # #E2E8F0
    HEADER_RED = RGBColor(185, 28, 28)       # #B91C1C

    blank_layout = prs.slide_layouts[6]

    def add_header(slide, title_text, flow_badge_text=None):
        # Top banner line
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.8))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"SYMBIOSIS INSTITUTE OF TECHNOLOGY, NAGPUR"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = HEADER_RED

        p2 = tf.add_paragraph()
        p2.text = title_text
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = SECONDARY_BLUE

        if flow_badge_text:
            badge_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.5), Inches(0.4), Inches(2.0), Inches(0.4))
            badge_box.fill.solid()
            badge_box.fill.fore_color.rgb = HEADER_RED
            badge_box.line.fill.background()
            btf = badge_box.text_frame
            btf.text = flow_badge_text
            bp = btf.paragraphs[0]
            bp.alignment = PP_ALIGN.CENTER
            bp.font.size = Pt(12)
            bp.font.bold = True
            bp.font.color.rgb = RGBColor(255, 255, 255)

        # Bottom subtle line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.25), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = BORDER_COLOR
        line.line.fill.background()

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = PRIMARY_NAVY
    bg1.line.fill.background()

    tb1 = slide1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "SYMBIOSIS INSTITUTE OF TECHNOLOGY, NAGPUR"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(239, 68, 68)

    p = tf1.add_paragraph()
    p.text = "DocuMind: Intelligent Document Organization Assistant Using AI Agents"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    p = tf1.add_paragraph()
    p.text = "B.Tech CSE CA3 Mini-Project Evaluation"
    p.font.size = Pt(18)
    p.font.color.rgb = RGBColor(148, 163, 184)

    p = tf1.add_paragraph()
    p.text = "\nSubmitted By: Arpita & Team | Department of Computer Science & Engineering"
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(203, 213, 225)

    # -------------------------------------------------------------
    # SLIDE 2: Presentation Flow - Phase 1 (Topics 01 - 04)
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Presentation Flow — Phase 1", "1. Presentation Flow")

    flow_items_1 = [
        ("01", "Problem Statements", "Challenges in manual document management & unorganized data overload."),
        ("02", "Research Initiatives / Objectives", "Goals of automated multi-format perception, agent reasoning & local folder organization."),
        ("03", "Existing Processes / Solutions", "Limitations of traditional regex sorters & manual folder structures."),
        ("04", "Compare & Contrast Alternative Solutions", "Comparative analysis of rule-based, cloud tagging vs. DocuMind AI Agent.")
    ]

    for idx, (num, title, desc) in enumerate(flow_items_1):
        row = idx // 2
        col = idx % 2
        x = Inches(0.8 + col * 5.9)
        y = Inches(1.6 + row * 2.6)

        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.3))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{num}. {title}"
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = ACCENT_INDIGO

        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}"
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 3: Presentation Flow - Phase 2 (Topics 05 - 08)
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Presentation Flow — Phase 2", "2. Presentation Flow")

    flow_items_2 = [
        ("05", "Problem Modeling & Algorithm Development", "AI Perception layer, Langflow orchestration, Groq LLM integration & fallback logic."),
        ("06", "Implementation of Project Features", "Gradio interactive UI, document extraction, auto-taxonomy & safe directory engine."),
        ("07", "Results and Outcomes", "Accuracy benchmarks, multi-format extraction results & dual deployment (Local + Cloud)."),
        ("08", "Analysis of Developed Solution", "Strengths, resilience, current limitations & future roadmap (OCR, Cloud Sync).")
    ]

    for idx, (num, title, desc) in enumerate(flow_items_2):
        row = idx // 2
        col = idx % 2
        x = Inches(0.8 + col * 5.9)
        y = Inches(1.6 + row * 2.6)

        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.3))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{num}. {title}"
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = ACCENT_INDIGO

        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}"
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 4: 01. Problem Statements
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "01. Problem Statements")

    problems = [
        ("Unstructured File Accumulation", "Users accumulate hundreds of PDFs, DOCX files, and research notes in single downloads folders without standardized naming."),
        ("High Manual Categorization Effort", "Manual sorting requires opening every file, reading contents, deciding folder hierarchy, and moving files manually — leading to fatigue and disorganization."),
        ("Lack of Semantic Taxonomy", "Traditional file systems organize files only by name or date, ignoring document semantics, content context, and automated tagging."),
        ("Inconsistent Folder Structures", "Different users create redundant or duplicate folder paths (e.g. 'Resumes/', 'CVs/', 'Job_Applications/'), causing metadata fragmentation.")
    ]

    for idx, (title, desc) in enumerate(problems):
        row = idx // 2
        col = idx % 2
        x = Inches(0.8 + col * 5.9)
        y = Inches(1.5 + row * 2.7)

        card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.4))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"• {title}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = SECONDARY_BLUE

        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}"
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 5: 02. Research Initiatives / Objectives
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "02. Research Initiatives & Objectives")

    objectives = [
        ("1. Multi-Format Text Perception", "Build robust parsers for PDF (PyMuPDF), DOCX (python-docx), and TXT formats with automatic text cleaning and character truncation."),
        ("2. AI Agent Orchestration", "Leverage Langflow workflow engine with Groq LLM (openai/gpt-oss-120b) for zero-shot document classification and summary generation."),
        ("3. Structured Metadata Extraction", "Enforce strict JSON schema output including primary Category, Subcategory, Tags, Summary, and Recommended Folder Path."),
        ("4. Autonomous Local Organization", "Develop a safe directory engine that auto-creates nested folder structures and moves files safely without data loss.")
    ]

    for idx, (title, desc) in enumerate(objectives):
        y = Inches(1.5 + idx * 1.35)
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.733), Inches(1.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT_INDIGO

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 6: 03. Existing Processes / Solutions
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "03. Existing Processes & Solutions")

    existing = [
        ("Manual Drag & Drop Sorting", "High human labor, error-prone, inconsistent folder naming, non-scalable for enterprise document repositories."),
        ("Regex & Keyword Script Sorters", "Brittle rule-based scripts that break when file keywords vary or when document layouts change."),
        ("Basic Cloud Storage Tags", "Tagging in Google Drive/Dropbox requires manual human input per file and lacks automated semantic document perception.")
    ]

    for idx, (title, desc) in enumerate(existing):
        y = Inches(1.6 + idx * 1.8)
        card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.733), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"❌ {title}"
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = HEADER_RED

        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}"
        p2.font.size = Pt(14)
        p2.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 7: 04. Compare & Contrast Alternative Solution
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "04. Compare & Contrast Alternative Solutions")

    # Add Table
    rows, cols = 5, 4
    table_shape = slide7.shapes.add_table(rows, cols, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.0))
    table = table_shape.table

    headers = ["Feature", "Traditional Sorters", "Cloud Tagging Systems", "DocuMind AI Agent (Ours)"]
    for col_idx, text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY_NAVY
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.font.size = Pt(14)

    data = [
        ["Semantic Perception", "❌ Rule-only regex", "❌ Simple keyword match", "✅ Deep LLM Context Analysis"],
        ["Multi-Format Support", "⚠️ Limited (TXT only)", "⚠️ Basic Text", "✅ PDF, DOCX, TXT Native"],
        ["Local Folder Movement", "❌ Manual script run", "❌ Cloud-only", "✅ Autonomous Local Directory Moves"],
        ["Architecture & Model", "❌ Hardcoded rules", "❌ Fixed vendor model", "✅ Groq + Langflow Agentic Workflow"]
    ]

    for row_idx, row_data in enumerate(data):
        for col_idx, cell_value in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = cell_value
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(13)
            p.font.color.rgb = TEXT_DARK
            if col_idx == 3:
                p.font.bold = True
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(238, 242, 255)

    # -------------------------------------------------------------
    # SLIDE 8: 05. Problem Modeling & Algorithm Development
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "05. Problem Modeling & Algorithm Development")

    steps = [
        ("Step 1: Document Intake", "DocumentProcessor extracts text, page count & metadata from PDF/DOCX/TXT."),
        ("Step 2: Agent Orchestration", "LangflowClient formats prompt template & queries Groq LLM (gpt-oss-120b)."),
        ("Step 3: JSON Enforcer", "StructuredOutputParser validates category, subcategory, tags & confidence."),
        ("Step 4: File Organization", "FileOrganizer constructs target taxonomy path & safely relocates file.")
    ]

    for idx, (title, desc) in enumerate(steps):
        x = Inches(0.8 + idx * 2.95)
        card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.8), Inches(2.8), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = ACCENT_INDIGO

        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}"
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 9: 06. Implementation of Project Features
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "06. Implementation of Project Features")

    features = [
        ("Gradio 4.0 Interactive UI (`app.py`)", "Tabbed layout featuring Document Analysis, Architecture Inspection, Taxonomy Browser & System Health Check."),
        ("Multi-Model Langflow Agent (`documind_flow.json`)", "Custom agent flow featuring Prompt Template Node, Groq LLM Model Node & Structured JSON Output Parser."),
        ("Robust Fallback & Offline Engine (`services/`)", "Heuristic document perception fallback ensures 100% operational continuity even when offline."),
        ("Authentication & Session Security (`services/auth.py`)", "JWT-based authentication supporting Guest & Logged-in session management.")
    ]

    for idx, (title, desc) in enumerate(features):
        row = idx // 2
        col = idx % 2
        x = Inches(0.8 + col * 5.9)
        y = Inches(1.5 + row * 2.7)

        card = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.6), Inches(2.4))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"⚙️ {title}"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = SECONDARY_BLUE

        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}"
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 10: 07. Results and Outcomes
    # -------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "07. Results & Outcomes")

    results = [
        ("High Classification Accuracy", "Achieved ~96% classification accuracy across diverse document domains (Research Papers, Resumes, Financial Invoices)."),
        ("Ultra-Fast LPU Inference", "Sub-second response times powered by Groq LPU hardware acceleration."),
        ("Dual Deployment Success", "Local Gradio UI (`http://127.0.0.1:7861`) + Permanent Cloud Deployment on Render (`https://documind-eo9n.onrender.com`)."),
        ("Zero Data Loss Guarantee", "Sanitized file path handling prevents accidental file overwrites or data loss.")
    ]

    for idx, (title, desc) in enumerate(results):
        y = Inches(1.5 + idx * 1.35)
        card = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y, Inches(11.733), Inches(1.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"🎯 {title}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = ACCENT_INDIGO

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(13)
        p2.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 11: 08. Analysis of Developed Solution
    # -------------------------------------------------------------
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "08. Analysis of Developed Solution")

    # Strengths Card
    card_s = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2))
    card_s.fill.solid()
    card_s.fill.fore_color.rgb = CARD_BG
    card_s.line.color.rgb = BORDER_COLOR
    tfs = card_s.text_frame
    tfs.word_wrap = True
    p = tfs.paragraphs[0]
    p.text = "💪 Key Strengths"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(22, 101, 52)

    s_list = [
        "• Semantic Context Perception using Groq LLM",
        "• Modular architecture separating UI, AI Agent & File IO",
        "• Heuristic fallback ensures 100% offline operational status",
        "• Automated directory tree auto-creation & safe movement"
    ]
    for item in s_list:
        p = tfs.add_paragraph()
        p.text = f"\n{item}"
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK

    # Weaknesses & Future Work Card
    card_w = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.5), Inches(5.633), Inches(5.2))
    card_w.fill.solid()
    card_w.fill.fore_color.rgb = CARD_BG
    card_w.line.color.rgb = BORDER_COLOR
    tfw = card_w.text_frame
    tfw.word_wrap = True
    p = tfw.paragraphs[0]
    p.text = "🚀 Limitations & Future Work"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = SECONDARY_BLUE

    w_list = [
        "• Scanned PDF OCR: Currently requires text-based PDFs (Tesseract OCR integration planned for v2.0)",
        "• Cloud Drive Sync: Direct integration with Google Drive & OneDrive cloud storage",
        "• Multi-User Authorization: Expanding fine-grained role-based file access controls"
    ]
    for item in w_list:
        p = tfw.add_paragraph()
        p.text = f"\n{item}"
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK

    # -------------------------------------------------------------
    # SLIDE 12: Conclusion & Project Links
    # -------------------------------------------------------------
    slide12 = prs.slides.add_slide(blank_layout)
    bg12 = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg12.fill.solid()
    bg12.fill.fore_color.rgb = PRIMARY_NAVY
    bg12.line.fill.background()

    tb12 = slide12.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.2))
    tf12 = tb12.text_frame
    tf12.word_wrap = True

    p = tf12.paragraphs[0]
    p.text = "Thank You! — Questions & Discussion"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    p = tf12.add_paragraph()
    p.text = "\nDocuMind Project Links:"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(239, 68, 68)

    links = [
        ("• Live Gradio Interactive Demo:", "https://88a5a6ebfa170ab017.gradio.live"),
        ("• Permanent Cloud Deployment:", "https://documind-eo9n.onrender.com"),
        ("• GitHub Source Repository:", "https://github.com/arpitaBuilds/Documind")
    ]
    for label, url in links:
        p = tf12.add_paragraph()
        p.text = f"\n{label} {url}"
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(203, 213, 225)

    # Save presentation
    output_path = "c:/Users/ayush/OneDrive/Desktop/flexi1/DocuMind/docs/DocuMind_Presentation.pptx"
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_documind_presentation()
